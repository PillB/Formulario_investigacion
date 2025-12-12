from pathlib import Path

import pytest
pytest.importorskip("docx")
from docx import Document

import report_builder
from tests.test_report_builder import sample_case_data


@pytest.mark.skipif(
    not report_builder.DOCX_AVAILABLE, reason=report_builder.DOCX_MISSING_MESSAGE
)
def test_mandatory_header_structure_preserved(tmp_path: Path, sample_case_data):
    output = tmp_path / "structure.docx"

    report_builder.build_docx(sample_case_data, output)

    document = Document(output)
    header_table = document.tables[0]

    assert len(header_table.rows) == 11
    assert all(len(row.cells) == 4 for row in header_table.rows)

    labels = [header_table.cell(idx, 0).text for idx in range(len(header_table.rows))]
    assert labels == [
        "Dirigido a",
        "Referencia",
        "Área de Reporte",
        "Categoría del evento",
        "Importe investigado",
        "Pérdida total",
        "Vencido",
        "Castigo",
        "Analítica Contable",
        "Producto",
        "N° de Reclamos",
    ]

    # Valores de columnas combinadas deben conservar el texto original
    assert header_table.cell(0, 1).text
    assert header_table.cell(1, 1).text == header_table.cell(1, 3).text
